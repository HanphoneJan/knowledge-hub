"""文档解析/结构修复（D167/D168/D169）的测试。

- PDF 分栏阅读顺序（`order_lines`）
- 表格 Markdown 化（`grid_to_markdown`）
- DOCX 标题层级、XLSX 列对齐、CSV/HTML
"""

from __future__ import annotations

import pymupdf
from docx import Document as Docx
from openpyxl import Workbook

from services.ingest.chunker import Chunk, build_payload, chunks_from_parsed
from services.ingest.locations import (
    ParsedChunk,
    describe_table,
    grid_to_json,
    grid_to_markdown,
    order_lines,
    table_text,
)
from services.ingest.parsers import parse_to_chunks
from services.retrieval.ask import _evidence_text

# ============================================================
# PDF 分栏阅读顺序
# ============================================================


def _line(text: str, x0: float, y0: float, x1: float) -> dict:
    return {"text": text, "x0": x0, "y0": y0, "y1": y0 + 12, "x1": x1,
            "bbox": (x0, y0, x1, y0 + 12)}


def test_order_lines_single_column_is_plain_yx_sort():
    """单栏必须退化为 (y, x) —— 与 D167 之前逐字相同（回归守卫）。"""
    lines = [_line(f"L{i}", 50, i * 20, 250) for i in range(10)]
    # 打乱输入顺序
    shuffled = [lines[3], lines[0], lines[9], lines[1], *lines[2:9:1]]
    assert order_lines(shuffled) == sorted(
        shuffled, key=lambda d: (round(d["y0"], 1), d["x0"])
    )


def test_order_lines_two_columns_reads_left_then_right():
    """双栏：同一 y 上左右栏的行不能相邻 —— 必须左栏读完再读右栏。"""
    left = [_line(f"L{i}", 50, i * 20, 250) for i in range(5)]
    right = [_line(f"R{i}", 320, i * 20, 520) for i in range(5)]
    # 交错输入（模拟底层对象的存储顺序）
    interleaved = [x for pair in zip(left, right) for x in pair]
    out = order_lines(interleaved)
    texts = [l["text"] for l in out]
    assert texts == [f"L{i}" for i in range(5)] + [f"R{i}" for i in range(5)]


# ============================================================
# 表格 Markdown
# ============================================================


def test_grid_to_markdown_basic():
    md = grid_to_markdown([["名称", "数量"], ["苹果", "3"]])
    assert "| 名称 | 数量 |" in md
    assert "| 苹果 | 3 |" in md


def test_grid_to_markdown_escapes_pipe_and_pads():
    md = grid_to_markdown([["a|b", "c"], ["1", "2", "3"]])
    assert "a\\|b" in md, "单元格里的竖线要转义，否则表格结构被破坏"
    # 每行都是合法的 markdown 表行（首尾都是竖线）
    for line in md.splitlines():
        assert line.startswith("|") and line.endswith("|")


# ============================================================
# 表格 JSON 化 + 表格描述（D203）
# ============================================================


def test_grid_to_json_keeps_structure_and_pads():
    j = grid_to_json([["姓名", "年龄"], ["张三", 20], ["李四"]])
    assert j is not None
    assert j["header"] == ["姓名", "年龄"], "首行当表头"
    assert j["rows"] == [["张三", "20"], ["李四", ""]], "短行补空、值转成字符串"
    assert j["row_count"] == 3 and j["col_count"] == 2


def test_grid_to_json_does_not_escape_pipe():
    """⚠️ 与 `grid_to_markdown` 的关键区别：JSON 里 `|` **原样保留**。

    转义是渲染层的事；把 `a\\|b` 写进数据会让程序化读到的值和原文不一致。
    """
    j = grid_to_json([["a|b"]])
    assert j is not None and j["header"] == ["a|b"]


def test_grid_to_json_empty_and_blank_rows():
    assert grid_to_json([]) is None
    assert grid_to_json([["", " "], ["", ""]]) is None, "整表皆空 → None（不是空对象）"


def test_describe_table_names_scale_and_header():
    desc = describe_table([["姓名", "年龄", "城市"], ["张三", "20", "北京"]])
    assert "3 列" in desc and "2 行" in desc
    assert "姓名" in desc and "年龄" in desc and "城市" in desc, "表头名要进描述"


def test_describe_table_empty_is_empty_string():
    assert describe_table([]) == ""
    assert describe_table([["", ""]]) == ""


def test_table_text_prepends_description_before_markdown():
    """表格 chunk 的文本 = 描述行 + Markdown 表（描述必须在前，便于检索）。"""
    text = table_text([["a", "b"], ["1", "2"]])
    assert text.startswith("表格描述：")
    assert "| a | b |" in text
    assert text.index("表格描述") < text.index("| a | b |")


def test_chunks_from_parsed_carries_table_json():
    """⚠️ ParsedChunk → chunker.Chunk 必须带上 `table_json`（漏接会静默丢结构）。"""
    parsed = [
        ParsedChunk(
            text=table_text([["a", "b"], ["1", "2"]]),
            kind="table",
            table_json=grid_to_json([["a", "b"], ["1", "2"]]),
        )
    ]
    chunks = chunks_from_parsed(parsed)
    assert chunks[0].table_json == {"header": ["a", "b"], "rows": [["1", "2"]],
                                    "row_count": 2, "col_count": 2}


# ============================================================
# CSV
# ============================================================


def test_parse_csv_keeps_header_and_columns(tmp_path):
    p = tmp_path / "t.csv"
    p.write_text("姓名,年龄\n张三,20\n李四,21\n", encoding="utf-8")
    chunks, _, _ = parse_to_chunks(p, file_name="t.csv")
    assert chunks and chunks[0].kind == "table"
    text = "\n".join(c.text for c in chunks)
    assert "姓名" in text and "张三" in text
    assert "张三 | 20" in text, "列要对齐（表头 2 列，值也 2 列）"
    # D203：结构化网格也要落地（供前端渲染真表格）
    assert chunks[0].table_json is not None
    assert chunks[0].table_json["header"] == ["姓名", "年龄"]
    assert ["张三", "20"] in chunks[0].table_json["rows"]


# ============================================================
# XLSX
# ============================================================


def test_parse_xlsx_aligns_columns_and_names_sheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["名称", None, "年龄"])      # 中间空表头
    ws.append(["张三", "备注", "20"])
    p = tmp_path / "t.xlsx"
    wb.save(p)

    chunks, _, _ = parse_to_chunks(p, file_name="t.xlsx")
    text = "\n".join(c.text for c in chunks)
    assert "[工作表: 数据]" in text, "多 sheet 文档要能分清来源"
    assert "张三 | 备注 | 20" in text, "空表头不能把后面的列挤错位"
    # D203：结构化网格保留原单元格（空表头是空串，不是被丢掉）
    assert chunks[0].table_json is not None
    assert chunks[0].table_json["header"] == ["名称", "", "年龄"]
    assert chunks[0].table_json["rows"] == [["张三", "备注", "20"]]


# ============================================================
# DOCX 标题层级
# ============================================================


def test_parse_docx_preserves_heading_levels(tmp_path):
    doc = Docx()
    doc.add_heading("第一章", level=1)
    doc.add_paragraph("这是第一章的正文内容，足够长以形成一个块。" * 3)
    doc.add_heading("小节", level=2)
    doc.add_paragraph("这是小节里的正文内容。" * 3)
    p = tmp_path / "t.docx"
    doc.save(p)

    chunks, _, _ = parse_to_chunks(p, file_name="t.docx")
    texts = "\n".join(c.text for c in chunks)
    assert "# 第一章" in texts, "H1 应还原成 markdown 标题"
    paths = {c.section_path for c in chunks}
    assert "第一章" in paths
    assert any(sp and "小节" in sp for sp in paths), "H2 的路径应带上级"


# ============================================================
# HTML
# ============================================================


def test_parse_html_uses_headings(tmp_path):
    p = tmp_path / "t.html"
    p.write_text(
        "<html><body><h2>标题A</h2><p>这是一段正文内容。</p>"
        "<script>var x=1;</script></body></html>",
        encoding="utf-8",
    )
    chunks, _, _ = parse_to_chunks(p, file_name="t.html")
    text = "\n".join(c.text for c in chunks)
    assert "## 标题A" in text
    assert "var x" not in text, "脚本内容不能进索引"


# ============================================================
# Parent-Child 元数据接线（D165）
# ============================================================


def test_build_payload_includes_parent_and_section():
    c = Chunk(text="x", parent_text="父块上下文", section_path="H1 > H2")
    pl = build_payload(
        c, chunk_index=0, kb_id=1, doc_id=1, workspace_id=1,
        visibility="private", source_kind="file", owner_id=1,
        file_name="f.md", mime_type=None,
    )
    assert pl["parent_text"] == "父块上下文"
    assert pl["section_path"] == "H1 > H2"


def test_evidence_text_prefers_parent_context():
    class WithParent:
        payload = {"parent_text": "父块上下文"}
        text = "子块正文"

    class WithoutParent:
        payload = {}
        text = "子块正文"

    assert _evidence_text(WithParent()) == "父块上下文"
    assert _evidence_text(WithoutParent()) == "子块正文"
