"""🔴 D131：Office 文档的 L3 坐标（从**转换版**抽，而不是迁移原件坐标）。

## 这组测试纠正的是**问题的前提**

CONVENTIONS §6.10 写的是"Office 文档坐标：走 LibreOffice 转 PDF 后排版会变，
**需重新评估**"，`office.py` 的模块 docstring 写着"本模块不做的事：
把 L3 坐标从原件迁移到转换版 PDF"。

评估结果（`deploy/verify_l3_office.py`，真实 docx）：

| 事实 | 实测 |
|---|---|
| Office chunk 有坐标吗 | **0/8**（`_parse_docx` 只出文本，docstring 就写着"无坐标"）|
| 转换版里能定位到吗 | **平均覆盖 99%，8/8 ≥95%** |

⇒ "迁移原件坐标"没有东西可迁；正确的做法是**从转换版自己抽** ——
用户看到的页图就是那份 PDF，文本与坐标**同源**，
于是 D106 担心的"框指向错误的文字"这个风险**根本不存在**。

## 这一组钉住四件事

1. **定位器本身**：贪心最长匹配、覆盖率不够就**不给框**（宁可没有也别错）
2. **门禁**：带框预览对 Office 开放，但**逐 chunk** 判断（有坐标才给框；
   chunk 不存在 → 403/404，与"存在但没坐标" → 415 可区分）
3. **同源不变量**：画框必须用**调用方给的 `src`**（不许自己去找原件）
4. **整条链路**：造 docx → 转换 → 补坐标 → 端点画框（真跑，不调内部函数）

## ⚠️ 素材一律**自建**，不借 eval 脚本留下的东西

库里那份 `.docx` 是 `eval/api_1c_check.py` §15 上传的，而它结束时会清场
→ 测试会变成"看 eval 跑到哪一步"的函数（实测在基线末尾直接 skip，
而基线开头不 skip —— 同一个测试两次跑出不同结论）。
所以下面那条链路测试在**临时目录**里 python-docx 造一份、LibreOffice 转一份。
"""
from __future__ import annotations

import ast
import hashlib
from pathlib import Path

import pytest

# 🔑 一个模块级 loop（CONVENTIONS.md §9「async 测试的坑」）：
#   `get_async_engine()` 是模块级单例，asyncpg 连接池绑定**创建它的那个 loop**，
#   而 anyio 的 pytest 插件**每个测试函数开新 loop** → 第二个用到 async 的测试
#   就报 `Event loop is closed`，且报错完全指不到真正的原因。
#   这与生产（FastAPI 单 loop）一致，也让连接池能跨用例复用。
_LOOP = __import__("asyncio").new_event_loop()


def _run(coro):  # noqa: ANN001, ANN201
    return _LOOP.run_until_complete(coro)


@__import__("pytest").fixture(scope="module", autouse=True)
def _release_async_engine():
    """退出本模块时**销毁 async 引擎单例**。

    🔑 不销毁的话，asyncpg 连接池会绑在本模块的 `_LOOP` 上，
    下一个用 async 的测试文件（`tests/test_retrieval_e2e.py` 有它自己的 loop）
    就会报 `got Future ... attached to a different loop` —— 而报错完全指不到
    真正的原因（看起来像并发问题）。实测踩过一次（§5.129 的同款"报错指错地方"）。
    """
    yield
    from services.common.db import dispose_engines

    _LOOP.run_until_complete(_close_pool())
    dispose_engines()


async def _close_pool() -> None:
    """把 async 引擎**优雅关掉**。

    ⚠️ 必须用 `await engine.dispose()`（async 版），不能用同步的
    `engine.sync_engine.pool.dispose()` —— 后者会留下一条
    `_terminate_graceful_close` 的 pending task，下次有人 await 它时就报
    `got Future ... attached to a different loop`。
    """
    from services.common.db import get_async_engine

    await get_async_engine().dispose()


ROOT = Path(__file__).resolve().parents[1]
DOCS_ROUTER = ROOT / "services" / "api" / "routers" / "documents.py"
HANDLERS = ROOT / "services" / "worker" / "handlers" / "__init__.py"

#: 造出来的 docx 的三段（每段够长，够定位器吃）
OFFICE_PARAS = (
    "我志愿加入本计划，为共同的目标努力、长期投入，"
    "在任何时候都坚持党的利益高于个人的利益。",
    "本人于二〇一九年九月入读本校计算机科学与技术专业，"
    "四年间成绩优秀，多次获得校级奖学金，并参与两项科研课题。",
    "担任班长期间，组织班级开展志愿服务六十余次，"
    "带领同学走访社区获得一致好评，本人也被评为优秀学生干部。",
)
OFFICE_MIME = ("application/vnd.openxmlformats-officedocument."
               "wordprocessingml.document")


def _lines(*rows):  # noqa: ANN001, ANN201
    """`LineBox` 列表的简写：`(text, bbox, page)`。"""
    from services.ingest.relocate import LineBox

    return [
        LineBox(text="".join(t.split()), bbox=bbox, page=page)   # type: ignore[arg-type]
        for t, bbox, page in rows
    ]


def _fn(path: Path, name: str) -> ast.AST:
    """按名字找函数（同步或异步都行 —— handler 是同步的、端点是异步的）。"""
    tree = ast.parse(path.read_text(encoding="utf-8"))
    return next(n for n in ast.walk(tree)
                if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
                and n.name == name)


# ============================================================
# 1. 定位器（纯函数）
# ============================================================


def test_normalize_drops_all_whitespace():
    """🔑 比较前必须去掉**所有**空白 —— 转换会重新折行、重排空格。

    实测保留空白时那份 docx 的覆盖率从 99% 掉到 ~50%
    （同一份素材，唯一差别就是这一条规范化）。
    """
    from services.ingest.relocate import normalize

    assert normalize("  你好\n\t世界  ") == "你好世界"
    assert normalize("") == ""


def test_greedy_match_finds_text_split_across_lines():
    """🔑 chunk 常常横跨多行/多段，所以要**逐片段**定位并合成多个框。

    第一版用 `page.search_for(chunk 前 16 字)`，实测命中率只有 50% ——
    它量的是 needle 的失败，不是定位能力的上限（§5.126 的同款：
    量出来的数字回答的不是被问的那个问题）。
    """
    from services.ingest.relocate import relocate_text

    lines = _lines(
        ("今天我第一次提交入会申请", (0.1, 0.1, 0.9, 0.12), 1),
        ("心里既紧张又激动", (0.1, 0.13, 0.9, 0.15), 2),
    )
    cov, boxes = relocate_text("今天我第一次提交入会申请\n心里既紧张又激动", lines)
    assert cov == 1.0, f"覆盖率 {cov}"
    assert [b.page for b in boxes] == [1, 2], boxes


def test_unlocatable_text_scores_zero_and_yields_no_boxes():
    from services.ingest.relocate import relocate_text

    lines = _lines(("完全无关的一句话", (0.1, 0.1, 0.9, 0.12), 1))
    cov, boxes = relocate_text("这是一段转换版里根本没有的文字", lines)
    assert cov == 0.0 and boxes == []


def test_low_coverage_yields_no_locations_at_all():
    """🔑 覆盖率不够 → **没有这个键**（而不是空数组）。

    ⚠️ 写空数组会让前端以为"有坐标但这次没画"，那是另一种误导；
    而"没有键"= 这份 chunk 没有坐标 → 退化成整页/无框（D60）。
    """
    from services.ingest.relocate import MIN_COVERAGE, relocate_chunks

    lines = _lines(("开头能对上后面全对不上", (0.1, 0.1, 0.9, 0.12), 1))
    out = relocate_chunks([(0, "开头能对上后面全对不上" + "×" * 200)], lines)
    assert 0 not in out, "覆盖率不够却给了坐标"
    assert MIN_COVERAGE > 0.5


def test_short_fragments_are_not_matched():
    """⚠️ 短于 8 字的片段**不参与匹配** —— 满页都是"党的"，匹配到也不知道是哪句。"""
    from services.ingest.relocate import relocate_text

    lines = _lines(("党的", (0.1, 0.1, 0.2, 0.12), 1))
    cov, _ = relocate_text("党的", lines)
    assert cov == 0.0


def test_locations_shape_matches_the_contract():
    """🔑 写进 `chunks.locations` 的形状必须是 `{"page": int, "bbox": [4 个 0~1]}`。

    前端 `lib/boxes.ts` 与服务端画框都按这个形状读；
    写成别的形状症状是"框画不出来"而不是报错（§5.57）。
    """
    from services.ingest.relocate import relocate_chunks

    lines = _lines(("我志愿加入本计划", (0.1, 0.2, 0.9, 0.22), 3))
    out = relocate_chunks([(7, "我志愿加入本计划")], lines)
    loc = out[7][0]
    assert set(loc) == {"page", "bbox"}, loc
    assert loc["page"] == 3
    assert len(loc["bbox"]) == 4 and all(0.0 <= v <= 1.0 for v in loc["bbox"]), loc


def test_relocated_boxes_are_not_treated_as_table_cells():
    """🔑 补出来的框**不带** `row`/`col` —— 前端才不会把它当表格网格画。

    🔑 `web/src/lib/boxes.ts::isTableLocs` 的判据是"有没有 row/col"。
    一旦补上 row/col，Office 的表格 chunk（`row_count/col_count` 有值）
    就会被画成网格，而网格的行列数与行框数根本对不上 → **框指向错误的位置**，
    正是 D131 要避免的那类误导。
    """
    from services.ingest.relocate import relocate_chunks

    lines = _lines(("项目名称负责人张三丰同学", (0.1, 0.1, 0.9, 0.12), 1))
    out = relocate_chunks([(0, "项目名称负责人张三丰同学")], lines)
    for loc in out[0]:
        assert "row" not in loc and "col" not in loc, loc


# ============================================================
# 2. 门禁与不覆盖（接线）
# ============================================================


def test_preview_endpoint_no_longer_refuses_office_outright():
    """🔑 D106 的「带框预览只支持 PDF」改成**逐 chunk 判断**。

    判据落在**调用点**上（§5.112）：端点里必须出现
    「需要转换 → 取转换版 → 这个 chunk 有没有坐标」的分支，
    而不是一句"不支持"。
    """
    body = ast.unparse(_fn(DOCS_ROUTER, "get_preview_with_boxes"))
    assert "needs_conversion" in body, (
        "带框预览没有区分 PDF 与 Office（D131 的门禁丢了）"
    )
    assert "_chunk_boxes_state" in body, (
        "没有逐 chunk 判断坐标是否存在 —— 要么全开（会给错框），要么全关（功能不成立）"
    )
    # 🔑 三态：不存在 → 404/403，"存在但没坐标" → 415。
    #   第一版把两者都答成"没坐标"，打错索引也回 415，
    #   于是"预览坏了"与"你写错了索引"在客户端无法区分（§0.4 的形状）。
    assert "if state is None" in body and "ResourceNotFound" in body, (
        "chunk 不存在时必须回 404/403，而不是 415"
    )


def test_handler_fills_locations_after_conversion():
    """🔑 坐标是在 `KIND_PDF_PREVIEW`（转换成功之后）补的。

    放在这里而不是 `ingest` 里，是因为**转换产物那时还不存在**；
    而放在 API 层又会让"读"操作产生写副作用。
    """
    body = ast.unparse(_fn(HANDLERS, "handle_pdf_preview"))
    assert "_locate_office_chunks" in body, (
        "转换之后没有补坐标 —— 带框预览对 Office 就永远是关的"
    )
    # ⚠️ 判据是"挑『还没有坐标』的那个 helper"，**不是** `locations IS NULL` ——
    #   那一列是 JSONB，存进去的常常是 JSON `null`，SQL 侧一条都筛不出来
    #   （实测静默补了 0 个）。见 `locations.has_locations`。
    #   这里判**函数体**而不是整份文件：注释里恰好写着那个反面例子，
    #   整文件匹配会把自己告发（§5.113 的老坑）。
    body2 = ast.unparse(_fn(HANDLERS, "_locate_office_chunks"))
    assert "has_locations(" in body2, (
        "补坐标时没有只挑『还没有坐标』的 chunk —— 会覆盖 PDF 的原件坐标"
    )
    assert "locations.is_(None)" not in body2.replace(" ", ""), (
        "用 `locations IS NULL` 判有没有坐标 —— JSONB 存的是 JSON null，"
        "筛出来是空集（D131 实测踩过）"
    )


def test_relocation_failure_does_not_fail_the_job():
    """🔑 坐标是**增强**：抽取失败不许让 `pdf_preview` 任务失败。

    与"转换失败不动文档状态"同一条纪律（转换失败也只是标任务失败）——
    反过来做会让人以为"文件没导入成功"，而它照样能搜能引用。
    """
    body = ast.unparse(_fn(HANDLERS, "_locate_office_chunks"))
    assert "except Exception" in body or "except BaseException" in body, (
        "抽坐标失败会往外抛 —— 预览明明成功了却把任务标成失败"
    )
    assert "raise" not in body, "抽坐标失败会抛出去（应该只记日志）"


def test_box_rendering_uses_the_same_file_the_boxes_came_from():
    """🔑 画框**不许**自己去找原件 —— 必须用调用方传进来的 `src`。

    🔴 D131 的第一版就是在这里出的 500：`_render_optional_boxes` 沿用
    D106 的老规矩"强制用原件 `blob_path(doc.blob_sha)`"，
    给 Office 画框时打开的是 **docx** → pymupdf 报 `is no PDF` → 500。

    判据是**不变量**"坐标与被渲染的页来自同一个文件"：
    PDF 的坐标按原件算（D33，`src` = 原件）、Office 的坐标按转换版算
    （D131，`src` = `preview.pdf`）—— 两者都由调用方决定，
    这个函数只管画。所以它的函数体里**不该再出现 `blob_path`**。
    """
    body = ast.unparse(_fn(DOCS_ROUTER, "_render_optional_boxes"))
    assert "blob_path" not in body, (
        "画框时又自己去解析原件 —— 坐标与页图可能不同源"
        "（Office 的坐标在转换版上，原件是 docx → pymupdf 报 is no PDF）"
    )
    assert "render_page_with_boxes(src" in body.replace(" ", ""), (
        "画框没有用传进来的 src"
    )
    # 坐标缺失要走 404 而不是画一张无框的图（否则"框画不出来"没有任何提示）
    assert "has_locations" in body, (
        "没有判坐标是否真的存在（JSON null 会被当成有坐标）"
    )


def test_office_module_docstring_points_at_d131():
    """🔑 `office.py` 曾经写着"本模块不做的事：迁移坐标"。

    D131 的结论是"坐标该从转换版抽"，而**抽取发生在 worker handler 里**
    （不是这个模块）—— 所以这句话要改成指向 D131，而不是留着
    一条已经不成立的免责声明。
    """
    text = (ROOT / "services" / "ingest" / "office.py").read_text(encoding="utf-8")
    assert "D131" in text, (
        "office.py 还在说「本模块不做的事：把 L3 坐标迁移到转换版」"
        "—— D131 已经有结论了，这句话会误导下一个人"
    )


def test_has_locations_judges_content_not_sql_nullness():
    """🔑 `has_locations` 看**内容**：JSON `null` / `[]` / `[[]]` 都不算有坐标。

    ⚠️ 这是 D131 里被咬了三次的地方：评估脚本数出"8/8 有坐标"、
    补坐标的 handler 静默补了 0 个、API 门禁会放行一份没有框的 chunk ——
    三次都是因为 `locations` 是 JSONB 而 SQL 侧判不出来。
    """
    from services.ingest.locations import has_locations

    assert has_locations(None) is False            # SQL NULL
    assert has_locations([]) is False              # Python 空列表
    assert has_locations([[]]) is False            # 有列表但没有框
    assert has_locations([{"page": 1}]) is False    # 有框但缺 bbox
    assert has_locations({"page": 1}) is False      # 形状不对的历史脏数据
    assert has_locations([{"page": 1, "bbox": [0, 0, 1, 1]}]) is True


# ============================================================
# 3. 整条链路（自建素材，真跑）
# ============================================================


def _pg_alive() -> bool:
    from sqlalchemy import text

    from services.common.db import sync_session_scope

    try:
        with sync_session_scope() as s:
            s.execute(text("SELECT 1"))
        return True
    except Exception:  # noqa: BLE001
        return False


pg = pytest.mark.skipif(not _pg_alive(), reason="Postgres 不在线")


@pytest.fixture(scope="module")
def office_pair(tmp_path_factory):  # noqa: ANN001, ANN201
    """造一份 docx 并转成 PDF（临时目录，不留残渣）。"""
    from services.ingest.office import ConversionError, convert_to_pdf, soffice_available

    if soffice_available() is None:
        pytest.skip("本机没有 LibreOffice（环境不具备，不是功能缺陷）")
    import docx

    d = docx.Document()
    for para in OFFICE_PARAS:
        d.add_paragraph(para)
    tmp = tmp_path_factory.mktemp("d131")
    src = tmp / "申请书.docx"
    d.save(str(src))
    pdf = tmp / "preview.pdf"
    try:
        convert_to_pdf(src, pdf)
    except ConversionError as exc:
        pytest.skip(f"LibreOffice 转换失败（环境问题）：{exc}")
    return src, pdf


@pg
def test_office_chunks_get_located_from_the_converted_pdf(office_pair):  # noqa: ANN001
    """🔑 真实排版下，chunk 文本在转换版上**大部分**能定位。

    ⚠️ 判据是**覆盖率**而不是"全部命中"：定位器对个别对不齐的字是容忍的
    （贪心跳过），而真实排版里总有几处差异（断行、标点）。
    但"一个都定位不到"必须红 —— 那说明素材或链路坏了。
    """
    from services.ingest.locations import parse_pdf_with_locations
    from services.ingest.relocate import LineBox, relocate_chunks

    _src, pdf = office_pair
    parsed = parse_pdf_with_locations(pdf)
    lines = [
        LineBox(
            text="".join((pc.text or "").split()),
            bbox=tuple(float(v) for v in loc["bbox"]),   # type: ignore[index]
            page=int(loc.get("page") or 0),
        )
        for pc in parsed.chunks
        for loc in (pc.locations or [])
        if loc.get("bbox")
    ]
    assert lines, "转换版一行都没抽到 —— 链路坏了"

    out = relocate_chunks(list(enumerate(OFFICE_PARAS)), lines)
    rate = len(out) / len(OFFICE_PARAS)
    assert rate >= 0.8, (
        f"只有 {len(out)}/{len(OFFICE_PARAS)} 段定位到了 —— "
        "定位器或转换链路坏了（真实排版实测 ≥95%）"
    )


@pg
def test_full_pipeline_from_upload_to_drawn_box(office_pair):  # noqa: ANN001
    """🔑 整条链路：解析 → 写库 → 补坐标 → **端点真的画出框**。

    🔑 为什么单测之外还要这一条：D131 真正的坑全在**接线**上 ——
    `locations.is_(None)` 筛出空集（JSON null，§5.127）、门禁放错位置、
    `_render_optional_boxes` 打开 docx → 500。这三条在"直接调纯函数"的
    单测里全都测不到，只有跑一遍链路才会露馅。

    ⚠️ 自己造文档并在 `finally` 里清掉（不留残渣给别的测试）；
    复用库里已有的库与用户（不新建用户 —— 那要过注册端点，徒增依赖）。
    """
    from sqlalchemy import select

    from services.api.routers.documents import get_preview_with_boxes
    from services.common.auth import UserContext
    from services.common.db import session_scope, sync_session_scope
    from services.common.models import (  # noqa: F401 — 注册表需要它们
        Blob, Chunk, Document, KnowledgeBase, User, Workspace,
    )
    from services.common.storage import blob_path
    from services.ingest.locations import has_locations
    from services.ingest.parsers import parse_to_chunks
    from services.worker.handlers import _locate_office_chunks

    src, pdf = office_pair
    data = src.read_bytes()
    sha = hashlib.sha256(data).hexdigest()

    with sync_session_scope() as s:
        user_id = s.execute(select(User.id).order_by(User.id).limit(1)).scalar_one_or_none()
        kb = s.execute(
            select(KnowledgeBase).where(KnowledgeBase.deleted_at.is_(None))
            .order_by(KnowledgeBase.id).limit(1)
        ).scalars().first()
        if user_id is None or kb is None:
            pytest.skip("库里还没有用户或知识库")
        kb_id = kb.id

    parsed_chunks, _scan, _img = parse_to_chunks(src, file_name="申请书.docx")
    assert parsed_chunks, "解析器没给出 chunk"

    blob_dir = blob_path(sha).parent
    blob_dir.mkdir(parents=True, exist_ok=True)
    blob_path(sha).write_bytes(data)
    pdf_dst = blob_dir / "preview.pdf"
    pdf_dst.write_bytes(pdf.read_bytes())

    with sync_session_scope() as s:
        # 🔑 `documents.blob_sha` 有外键指向 `blobs` —— 先有 blob 行，
        #   否则插入直接 IntegrityError（这次就撞到了）。
        s.merge(Blob(sha256=sha, size_bytes=len(data), mime_type=OFFICE_MIME,
                     storage_path=f"blobs/{sha[:2]}/{sha}/original", ref_count=1))
        s.flush()
        # ⚠️ 只给**必填**列（`created_at/updated_at` 有 default）——
        #   这个模型列很多，写全了反而容易在新增列时崩（这次就崩在
        #   `ingested_at` 上：那个列叫 `doc_source`）。
        doc = Document(
            kb_id=kb_id, file_name="d131-链路自测.docx", mime_type=OFFICE_MIME,
            size_bytes=len(data), blob_sha=sha, content_hash=sha,
            status="ready", status_msg="", chunk_count=len(parsed_chunks),
            uploaded_by=user_id, doc_source="upload",
        )
        s.add(doc)
        s.flush()
        for i, pc in enumerate(parsed_chunks):
            s.add(Chunk(
                doc_id=doc.id, chunk_index=i, content_norm=pc.text,
                # ⚠️ `page` 不是 `chunks` 的列（页码在 `locations[].page` 里）——
                #   第一次写就多给了它，TypeError。真实写入走
                #   `chunker.chunks_from_parsed`，这里只喂文本 + 坐标。
                locations=pc.locations or None, kind=pc.kind,
            ))
        s.flush()
        doc_id = doc.id

    try:
        with sync_session_scope() as s:
            n = _locate_office_chunks(s, doc_id, pdf_dst)
            s.flush()
            rows = s.execute(
                select(Chunk.chunk_index, Chunk.locations)
                .where(Chunk.doc_id == doc_id).order_by(Chunk.chunk_index)
            ).all()
        assert n > 0, "补坐标一个都没补上（JSON null 筛不出行？）"
        missing = [r[0] for r in rows if not has_locations(r[1])]
        assert not missing, f"这些 chunk 没拿到坐标：{missing}"

        ctx = UserContext(
            user_id=user_id, display_name="u", email="u@t.local", api_key_id=None,
            readable_kb_ids=frozenset({kb_id}), workspace_ids=frozenset(),
        )

        async def _render(chunk_index: int | None) -> tuple[int, bytes]:
            async with session_scope() as session:
                r = await get_preview_with_boxes(
                    doc_id=doc_id, ctx=ctx, session=session, page=1, width=600,
                    chunk_index=chunk_index,
                )
                return r.status_code, r.body

        plain_code, plain_img = _run(_render(None))
        boxed_code, boxed_img = _run(_render(0))
        # 🔴 回归守卫：普通页图（不带 chunk_index）对 Office 一直是能看的
        #   （D34 v1.1）。D131 第一版把门禁放在分支最前 → `?page=1` 也 415。
        assert plain_code == 200, f"普通页图对 Office 回归了：HTTP {plain_code}"
        assert boxed_code == 200, f"带框页图：HTTP {boxed_code}"
        assert len(boxed_img) != len(plain_img), (
            "带框图与无框图字节数完全相同 → 框可能没画"
        )
    finally:
        with sync_session_scope() as s:
            s.execute(Chunk.__table__.delete().where(Chunk.doc_id == doc_id))
            s.execute(Document.__table__.delete().where(Document.id == doc_id))
            # 引用计数归零才物理删 blob（D28 的规矩）：这里引用它的行已删
            s.execute(Blob.__table__.delete().where(Blob.sha256 == sha))
        for f in (blob_path(sha), pdf_dst):
            try:
                f.unlink()
            except OSError:
                pass


@pg
@pg
def test_docx_chunk_locations_shapes_all_agree_with_has_locations():
    """🔑 `locations` 里**每一条** docx chunk 的形状都合法，且与 `has_locations` 一致。

    ⚠️ 第一版断言的是"库里**存在** JSON `null` 的 docx chunk" ——
    那是**数据**的性质，不是代码的性质。D131 给 docx 补了坐标、eval 1c 又清掉了
    旧文档，于是它在**没有任何代码变化**的情况下自己红了（实测 2026-10-05，
    属 §5.76「断言的前提错了」那一族）。

    ⇒ 改成逐条断言两件事，它们都与"库里恰好有什么"无关：

    1. 形状只能是三种之一：SQL `NULL` / JSON `null` / 数组
       （出现第四种 ⇒ 所有 SQL 侧判据都要一起复核）
    2. **非数组 ⇒ `has_locations` 为 False**；数组 ⇒ 它与"含 bbox 的元素数 > 0"
       一致 —— 这才是 D131 那条不变式的真实内容（§5.127：JSONB 存 JSON `null`
       被咬了三次）
    """
    from sqlalchemy import select, text as sql_text

    from services.common.db import sync_session_scope
    from services.common.models import Chunk, Document
    from services.ingest.locations import has_locations

    with sync_session_scope() as s:
        doc_ids = s.execute(
            select(Document.id).where(Document.file_name.like("%.docx")).limit(50)
        ).scalars().all()
        if not doc_ids:
            pytest.skip("库里没有 .docx")
        shapes: dict[str, int] = {}
        mismatched: list[str] = []
        for did in doc_ids:
            for idx, kind, raw in s.execute(
                select(
                    Chunk.chunk_index,
                    sql_text("jsonb_typeof(locations)"),
                    Chunk.locations,
                ).where(Chunk.doc_id == did)
            ).all():
                kind = kind or "sql-null"
                shapes[kind] = shapes.get(kind, 0) + 1
                if kind not in ("null", "array", "sql-null"):
                    mismatched.append(f"chunk {did}#{idx}: 形状是 {kind}")
                    continue
                expect = kind == "array" and any(
                    isinstance(x, dict) and x.get("bbox") for x in (raw or [])
                )
                if has_locations(raw) is not expect:
                    mismatched.append(
                        f"chunk {did}#{idx}: {kind} → has_locations="
                        f"{has_locations(raw)}，应为 {expect}"
                    )
    assert mismatched == [], mismatched[:5]
    assert sum(shapes.values()) > 0, f"没有 docx chunk 行：{shapes}"

@pg
def test_has_locations_agrees_with_what_postgres_actually_round_trips():
    """🔑 写进去四种形状、**从 PG 读回来**，`has_locations` 都判对。

    🔴 这条才是 D131 那条不变式的**可执行**版本：坑不在 `has_locations` 的代码里，
    而在**PG 存回来的东西** —— 写 `None` 之后读回来可能是 SQL `NULL`、
    也可能是 JSON `null`（§5.127 说这个"被咬了三次"）。所以必须**真写真读**：
    只在内存里喂 `None` / `[]` **测不到**这个往返，而往返才是出事的地方。

    ⚠️ **不断言"写 `None` 回来一定是哪一种"** —— 那由 SQLAlchemy 与 psycopg 的
    版本决定，两种都合法。真正的性质是：**两种都判成"没有坐标"**，
    且都与 `jsonb_typeof` 报告的形状相容。

    🔑 **实测过的变异结果**（2026-10-05，写下来是因为结果与直觉相反）：

    | 变异 | 本条 | 结论 |
    |---|---|---|
    | `has_locations` 把"有数组"当"有坐标"（不看 `bbox`）| 🔴 **变红** | ✅ 判据咬得住 |
    | `if not raw` 改成 `if raw is None`（"只看 SQL NULL"）| 🟢 **仍然绿** | **等价变异**：后面 `isinstance(raw, dict)` 与 `any(...bbox)` 已经把空值判掉了 |

    ⚠️ 所以**别写"某个变异一定会让测试变红"这种话，除非实测过** ——
    它本身就是一条待验证的主张（本文件这次就是先写了、实测才发现是等价变异）。
    另外，只读真库里**已有**的行是**测不到**往返的：那些行恰好全是
    "有坐标"的数组，所以上面才要自己造探针行。
    """
    from sqlalchemy import select, text as sql_text

    from services.common.db import sync_session_scope
    from services.common.models import Chunk, Document
    from services.ingest.locations import has_locations

    cases = {
        # chunk_index → (写进去的 locations, 期望 has_locations)
        9101: (None, False),                                       # 空
        9102: ([], False),                                         # 空数组
        9103: ([{"page": 1}], False),          # 有元素但**没有 bbox**
        9104: ([{"page": 1, "bbox": [0, 0, 1, 1]}], True),         # 真有框
    }
    with sync_session_scope() as s:
        doc_id = s.execute(select(Document.id).limit(1)).scalar_one_or_none()
        if doc_id is None:
            pytest.skip("库里还没有任何文档（先跑一遍 eval/api_1c_check.py）")
        doc = s.get(Document, doc_id)
        made = [
            Chunk(doc_id=doc_id, chunk_index=idx, content_norm=f"d137-{idx}",
                  locations=val)
            for idx, (val, _) in cases.items()
        ]
        s.add_all(made)
        s.flush()
        probe_ids = [c.id for c in made]
        seen = s.execute(
            select(Chunk.chunk_index, Chunk.locations,
                   sql_text("jsonb_typeof(locations)"))
            .where(Chunk.id.in_(probe_ids))
        ).all()
        s.execute(Chunk.__table__.delete().where(Chunk.id.in_(probe_ids)))
        s.flush()

    read = {idx: (raw, kind) for idx, raw, kind in seen}
    assert set(read) == set(cases), f"没读回全部探针行：{sorted(read)}"
    wrong = [
        f"chunk_index={idx}: PG 说是 {kind}，has_locations={has_locations(raw)}，"
        f"应为 {cases[idx][1]}"
        for idx, (raw, kind) in read.items()
        if has_locations(raw) is not cases[idx][1]
    ]
    assert wrong == [], wrong
    # 形状只允许三种：SQL NULL / JSON null / array
    odd = [f"{idx}:{kind}" for idx, (_, kind) in read.items()
           if kind not in ("null", "array", "sql-null")]
    assert odd == [], f"出现了第四种形状：{odd}"
    # 🔑 顺带把 D131 的结论**钉成断言**：`None` 往返之后 `IS NULL` 与
    #   `has_locations` **不保证一致**（不一致的那一半就是 §5.127 的坑）。
    #   这里只断言"两者都被 has_locations 判成没有坐标"，不挑哪一种。
    assert not has_locations(read[9101][0]), "写 None 的那行被当成了有坐标"
